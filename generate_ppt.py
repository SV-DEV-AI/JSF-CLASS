import sys
import subprocess
import os

def install_and_import(package):
    try:
        import pptx
    except ImportError:
        print(f"Installing {package}...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", package])

install_and_import('python-pptx')

from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()

# Slide 1: Title Slide
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "Networking Lab Project Report"
slide.placeholders[1].text = "A Comprehensive Analysis of Enterprise Network Architectures\n\nExperiment 6: Smart Office Network Challenge\nExperiment 7: Bank ATM Network Simulation"

# Slide 2: Abstract/Overview
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Project Overview"
slide.placeholders[1].text = (
    "Demonstration of critical Layer 2 and Layer 3 networking protocols.\n\n"
    "1. Experiment 6: Smart Office Network\n"
    "   - VLAN segmentation, ROAS, NAT, and Zero-Trust ACLs.\n\n"
    "2. Experiment 7: Bank ATM Network\n"
    "   - Highly available WAN, redundant topology, and OSPF dynamic routing."
)

# Slide 3: Exp 6 Scenario
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Exp 6: Smart Office Scenario"
slide.placeholders[1].text = (
    "Aim: Design and secure a multi-department corporate network.\n\n"
    "Requirements:\n"
    "• Departments: HR, Accounts, Developers, Testing, Server Room.\n"
    "• Access Policy: Only Developers can access the Server Room.\n"
    "• Privacy: HR data isolated, restricted to outbound internet only.\n"
    "• Global Connectivity: Outbound internet access via NAT."
)

# Slide 4: Exp 6 Architecture
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Exp 6: Network Architecture"
slide.placeholders[1].text = (
    "• VLANs (Virtual LANs): Logical broadcast containment.\n"
    "• 802.1Q Trunking & ROAS: Efficient inter-VLAN routing.\n"
    "• NAT Overload (PAT): Maps internal private subnets to public IP.\n"
    "• Extended ACLs: Strict filtering based on Source/Dest IP & Protocol."
)

# Slide 5: Exp 6 Outcomes
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Exp 6: Testing & Results"
slide.placeholders[1].text = (
    "• Internet Connectivity: Successful pings to Google Server (8.8.8.8).\n"
    "• Security Validation: Unauthorized pings (HR/Accounts -> Server) successfully blocked by router firewall.\n"
    "• Authorized Access: Developer PC successfully connected to Server Room.\n"
    "• Result: Zero-trust policies completely fulfilled."
)

# Slide 6: Exp 7 Scenario
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Exp 7: Bank ATM Network"
slide.placeholders[1].text = (
    "Aim: Simulate a highly available, multi-city ATM WAN.\n\n"
    "Requirements:\n"
    "• Connect critical ATM servers in Patna, Delhi, and Mumbai.\n"
    "• Extreme fault tolerance and high availability.\n"
    "• Automatic rerouting during catastrophic physical link failures."
)

# Slide 7: Exp 7 Architecture
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Exp 7: High Availability Design"
slide.placeholders[1].text = (
    "• Topology: Redundant 3-city triangular WAN core.\n"
    "• Protocol: OSPF (Open Shortest Path First) Link-State Routing.\n"
    "• Mechanism: Dijkstra's SPF Algorithm calculates absolute shortest paths.\n"
    "• Backbone Area 0: Facilitates rapid sharing of Link State Advertisements (LSAs)."
)

# Slide 8: Exp 7 Outcomes
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Exp 7: Failover Simulation"
slide.placeholders[1].text = (
    "• Convergence Verified: OSPF neighbors established 'FULL' adjacency.\n"
    "• Failure Simulated: Direct Delhi-Mumbai WAN link intentionally severed.\n"
    "• Automatic Failover: Tracert utility confirmed immediate rerouting through Patna.\n"
    "• Result: Zero downtime ATM transaction connectivity maintained."
)

# Slide 9: Conclusion
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Conclusion & Learning Outcomes"
slide.placeholders[1].text = (
    "• Demonstrated mastery of Layer 2 VLANs and Layer 3 Routing.\n"
    "• Proven application of NAT and Extended ACLs for cybersecurity.\n"
    "• Advanced understanding of Enterprise WAN design.\n"
    "• Real-time implementation of OSPF fault-tolerance in mission-critical systems."
)

output_path = os.path.join(r"C:\Users\Sarthak\Desktop\CN PROJECT TP 6", "Networking_Project_Presentation.pptx")
prs.save(output_path)
print(f"Presentation generated successfully at {output_path}")
